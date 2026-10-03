"""
KISAN URJA — Yuva Yodha Energy Tech Hackathon 2026
10-Slide Competition-Ready PowerPoint Generator

All content sourced from the actual GitHub repository and README.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ──────────────────── COLORS ────────────────────
BG_DARK       = RGBColor(0x0F, 0x17, 0x2A)
BG_CARD       = RGBColor(0x1E, 0x29, 0x3B)
BG_CARD_LIGHT = RGBColor(0x33, 0x41, 0x55)
EMERALD       = RGBColor(0x10, 0xB9, 0x81)
EMERALD_LIGHT = RGBColor(0x34, 0xD3, 0x99)
EMERALD_DIM   = RGBColor(0x05, 0x96, 0x69)
AMBER         = RGBColor(0xF5, 0x9E, 0x0B)
AMBER_LIGHT   = RGBColor(0xFB, 0xBF, 0x24)
AMBER_DIM     = RGBColor(0xD9, 0x77, 0x06)
SKY           = RGBColor(0x38, 0xBD, 0xF8)
RED           = RGBColor(0xEF, 0x44, 0x44)
WHITE         = RGBColor(0xF1, 0xF5, 0xF9)
TEXT_SEC      = RGBColor(0x94, 0xA3, 0xB8)
TEXT_TER      = RGBColor(0x64, 0x74, 0x8B)
TRANSPARENT   = RGBColor(0x0F, 0x17, 0x2A)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

ASSETS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")

# ──────────────────── HELPERS ────────────────────

def set_slide_bg(slide, color=BG_DARK):
    """Set solid background color for a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                font_color=WHITE, bold=False, italic=False, alignment=PP_ALIGN.LEFT,
                font_name="Outfit"):
    """Add a textbox with a single run."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = font_color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font_name
    return txBox, tf, p


def add_rich_textbox(slide, left, top, width, height):
    """Add an empty textbox ready for rich text."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return txBox, tf


def add_run(paragraph, text, font_size=16, font_color=WHITE, bold=False,
            italic=False, font_name="Plus Jakarta Sans"):
    """Add a run to a paragraph."""
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = font_color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font_name
    return run


def add_paragraph(tf, text="", font_size=16, font_color=WHITE, bold=False,
                  italic=False, alignment=PP_ALIGN.LEFT, font_name="Plus Jakarta Sans",
                  space_before=0, space_after=0):
    """Add a new paragraph with a single run."""
    p = tf.add_paragraph()
    p.alignment = alignment
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if text:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.color.rgb = font_color
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font_name
    return p


def add_rect(slide, left, top, width, height, fill_color=BG_CARD, border_color=None, border_width=Pt(1), corner_radius=None):
    """Add a rounded rectangle shape."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    return shape


def add_bullet_list(slide, left, top, width, height, items, font_size=14,
                    text_color=TEXT_SEC, bullet_color=EMERALD, bold_prefix=True):
    """Add a bulleted list."""
    txBox, tf = add_rich_textbox(slide, left, top, width, height)

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(4)

        # Bullet character
        bullet_run = p.add_run()
        bullet_run.text = "● "
        bullet_run.font.size = Pt(8)
        bullet_run.font.color.rgb = bullet_color
        bullet_run.font.name = "Plus Jakarta Sans"

        # Check if item has bold prefix (text before "—")
        if bold_prefix and "—" in item:
            parts = item.split("—", 1)
            bold_r = p.add_run()
            bold_r.text = parts[0].strip()
            bold_r.font.size = Pt(font_size)
            bold_r.font.color.rgb = WHITE
            bold_r.font.bold = True
            bold_r.font.name = "Plus Jakarta Sans"

            rest_r = p.add_run()
            rest_r.text = " — " + parts[1].strip()
            rest_r.font.size = Pt(font_size)
            rest_r.font.color.rgb = text_color
            rest_r.font.name = "Plus Jakarta Sans"
        else:
            r = p.add_run()
            r.text = item
            r.font.size = Pt(font_size)
            r.font.color.rgb = text_color
            r.font.name = "Plus Jakarta Sans"

    return txBox


def add_slide_number(slide, num):
    """Add slide number in top-right."""
    add_textbox(slide, Inches(11.8), Inches(0.3), Inches(1.3), Inches(0.4),
                f"{num:02d} / 10", font_size=10, font_color=TEXT_TER,
                alignment=PP_ALIGN.RIGHT, font_name="JetBrains Mono")


def add_footer(slide, left_text="KISAN URJA — Yuva Yodha 2026",
               right_text="kisanurja.vercel.app"):
    """Add footer bar."""
    add_textbox(slide, Inches(0.8), Inches(6.9), Inches(5), Inches(0.4),
                left_text, font_size=9, font_color=TEXT_TER, font_name="Plus Jakarta Sans")
    add_textbox(slide, Inches(8), Inches(6.9), Inches(4.5), Inches(0.4),
                right_text, font_size=9, font_color=EMERALD,
                alignment=PP_ALIGN.RIGHT, font_name="JetBrains Mono")


def add_label(slide, left, top, text, color=EMERALD):
    """Add a small uppercase label."""
    add_textbox(slide, left, top, Inches(5), Inches(0.35),
                text, font_size=11, font_color=color, bold=True,
                font_name="Outfit")


def add_image_safe(slide, img_name, left, top, width=None, height=None):
    """Add image if it exists, skip silently if not."""
    path = os.path.join(ASSETS_DIR, img_name)
    if os.path.exists(path):
        kwargs = {"image_file": path, "left": left, "top": top}
        if width:
            kwargs["width"] = width
        if height:
            kwargs["height"] = height
        return slide.shapes.add_picture(**kwargs)
    return None


def add_divider_line(slide, left, top, width):
    """Add a thin gradient-like colored line."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = EMERALD
    shape.line.fill.background()
    return shape


# ──────────────────── BUILD PRESENTATION ────────────────────

def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]  # Blank layout

    # ═══════════════════ SLIDE 1: COVER ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    # Background image
    add_image_safe(slide, "cover_background.jpg", Inches(0), Inches(0),
                   width=SLIDE_W, height=SLIDE_H)

    # Semi-transparent overlay
    overlay = add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H,
                       fill_color=BG_DARK)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = BG_DARK
    # Set transparency via XML
    try:
        spPr = overlay._element.spPr
        solidFill_el = spPr.find(qn('a:solidFill'))
        if solidFill_el is not None:
            srgb = solidFill_el.find(qn('a:srgbClr'))
            if srgb is not None:
                from lxml import etree
                a = etree.SubElement(srgb, qn('a:alpha'))
                a.set('val', '45000')
    except Exception:
        pass  # non-fatal — overlay will just be opaque

    # Hackathon label
    add_textbox(slide, Inches(0), Inches(1.5), SLIDE_W, Inches(0.4),
                "YUVA YODHA ENERGY TECH HACKATHON 2026",
                font_size=12, font_color=EMERALD, bold=True,
                alignment=PP_ALIGN.CENTER, font_name="Outfit")

    # Title
    txBox, tf = add_rich_textbox(slide, Inches(1.5), Inches(2.2), Inches(10.3), Inches(1.3))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "KISAN ", font_size=56, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "URJA", font_size=56, font_color=AMBER, bold=True, font_name="Outfit")

    # Hindi subtitle
    add_textbox(slide, Inches(0), Inches(3.3), SLIDE_W, Inches(0.7),
                "किसान ऊर्जा", font_size=28, font_color=TEXT_SEC, bold=True,
                alignment=PP_ALIGN.CENTER, font_name="Noto Sans Devanagari")

    # Divider
    add_divider_line(slide, Inches(5.9), Inches(4.05), Inches(1.5))

    # Subtitle
    add_textbox(slide, Inches(1.5), Inches(4.3), Inches(10.3), Inches(0.6),
                "Autonomous Agricultural Intelligence & Solar Irrigation Platform",
                font_size=18, font_color=TEXT_SEC, bold=False,
                alignment=PP_ALIGN.CENTER, font_name="Outfit")

    # Positioning quote
    add_textbox(slide, Inches(1.5), Inches(5.0), Inches(10.3), Inches(0.5),
                '"Connecting crop water intelligence with solar-powered irrigation."',
                font_size=14, font_color=TEXT_TER, italic=True,
                alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")

    # Tags row
    tags = ["☀️ Solar-Synchronized", "🌾 FAO-56 Physics", "🛰️ Sentinel-2 Canopy"]
    tag_w = Inches(2.2)
    start_x = Inches(3.0)
    for i, tag_text in enumerate(tags):
        x = start_x + i * (tag_w + Inches(0.3))
        tag_shape = add_rect(slide, x, Inches(5.6), tag_w, Inches(0.4),
                             fill_color=RGBColor(0x14, 0x2E, 0x22),
                             border_color=EMERALD_DIM)
        add_textbox(slide, x, Inches(5.6), tag_w, Inches(0.4),
                    tag_text, font_size=10, font_color=EMERALD_LIGHT,
                    bold=True, alignment=PP_ALIGN.CENTER, font_name="Outfit")

    # Footer
    add_textbox(slide, Inches(0.8), Inches(6.85), Inches(5), Inches(0.35),
                "Challenge: Sustainable Agriculture", font_size=9,
                font_color=TEXT_TER, font_name="Plus Jakarta Sans")
    add_textbox(slide, Inches(8), Inches(6.85), Inches(4.5), Inches(0.35),
                "kisanurja.vercel.app", font_size=9, font_color=EMERALD,
                alignment=PP_ALIGN.RIGHT, font_name="JetBrains Mono")
    add_slide_number(slide, 1)


    # ═══════════════════ SLIDE 2: THE PROBLEM ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "THE PROBLEM")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(1.1))
    p = tf.paragraphs[0]
    add_run(p, "Farm irrigation is a ", font_size=32, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "water", font_size=32, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, " problem", font_size=32, font_color=WHITE, bold=True, font_name="Outfit")
    p2 = tf.add_paragraph()
    add_run(p2, "AND ", font_size=32, font_color=TEXT_SEC, font_name="Outfit")
    add_run(p2, "an ", font_size=32, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p2, "energy", font_size=32, font_color=AMBER, bold=True, font_name="Outfit")
    add_run(p2, " problem.", font_size=32, font_color=WHITE, bold=True, font_name="Outfit")

    # Left bullets
    bullets = [
        "When to irrigate? — Farmers must judge timing from weather, soil moisture, and crop stage—often separately.",
        "How much water? — Fixed schedules miss changing field conditions, risking over- or under-irrigation.",
        "Energy cost — Pumping consumes electricity or diesel. Tariff timing and solar availability are rarely factored in.",
        "Fragmented signals — Weather, soil, crop, water, and energy are treated as disconnected datasets.",
        "No unified recommendation — Farmers need one actionable decision, not five separate dashboards."
    ]
    add_bullet_list(slide, Inches(0.8), Inches(2.3), Inches(6.2), Inches(4.2), bullets, font_size=13)

    # Right: Fragmented icons
    icons = [("🌦️", "Weather"), ("🪨", "Soil"), ("🌾", "Crop"), ("💧", "Water"), ("⚡", "Energy")]
    row_y = Inches(2.6)
    for i, (icon, label) in enumerate(icons):
        x = Inches(7.8) + (i % 3) * Inches(1.5)
        y = row_y + (i // 3) * Inches(1.3)
        card = add_rect(slide, x, y, Inches(1.2), Inches(1.0),
                        fill_color=BG_CARD, border_color=TEXT_TER)
        add_textbox(slide, x, y + Inches(0.1), Inches(1.2), Inches(0.45),
                    icon, font_size=22, alignment=PP_ALIGN.CENTER)
        add_textbox(slide, x, y + Inches(0.6), Inches(1.2), Inches(0.3),
                    label, font_size=9, font_color=TEXT_TER,
                    alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")

    # Arrow down
    add_textbox(slide, Inches(7.8), Inches(5.0), Inches(4.5), Inches(0.4),
                "↓  fragmented decisions  ↓", font_size=12, font_color=TEXT_TER,
                alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")

    # Unified card
    unified = add_rect(slide, Inches(7.8), Inches(5.5), Inches(4.5), Inches(0.65),
                        fill_color=RGBColor(0x0D, 0x2B, 0x1F),
                        border_color=EMERALD)
    add_textbox(slide, Inches(7.8), Inches(5.55), Inches(4.5), Inches(0.55),
                "KISAN URJA brings these signals together.",
                font_size=13, font_color=EMERALD_LIGHT, bold=True,
                alignment=PP_ALIGN.CENTER, font_name="Outfit")

    add_footer(slide)
    add_slide_number(slide, 2)


    # ═══════════════════ SLIDE 3: OUR SOLUTION ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "OUR SOLUTION")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "One intelligence layer for ", font_size=30, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "water", font_size=30, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, ", ", font_size=30, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "crops", font_size=30, font_color=AMBER, bold=True, font_name="Outfit")
    add_run(p, " and ", font_size=30, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "energy", font_size=30, font_color=AMBER, bold=True, font_name="Outfit")
    add_run(p, ".", font_size=30, font_color=WHITE, bold=True, font_name="Outfit")

    # Three columns: INPUTS | INTELLIGENCE | OUTPUTS
    col_w = Inches(3.5)
    col_gap = Inches(0.4)

    # INPUTS column
    inp_x = Inches(0.8)
    add_textbox(slide, inp_x, Inches(1.8), col_w, Inches(0.35),
                "INPUTS", font_size=12, font_color=TEXT_TER, bold=True,
                font_name="Outfit")

    inputs = [
        ("🌤️", "Microclimate — Open-Meteo weather"),
        ("🪨", "Soil properties — ISRIC SoilGrids 250m"),
        ("🛰️", "Satellite — Copernicus Sentinel-2 (10m)"),
        ("🌾", "Crop/field — Species, stage, rooting depth"),
        ("☀️", "Solar/pump — PV capacity, pump HP, tariff"),
    ]
    for i, (icon, text) in enumerate(inputs):
        y = Inches(2.2) + i * Inches(0.85)
        add_rect(slide, inp_x, y, col_w, Inches(0.7), fill_color=BG_CARD, border_color=RGBColor(0x2D, 0x3F, 0x55))
        add_textbox(slide, inp_x + Inches(0.15), y + Inches(0.12), col_w - Inches(0.3), Inches(0.5),
                    f"{icon}  {text}", font_size=11, font_color=TEXT_SEC, font_name="Plus Jakarta Sans")

    # INTELLIGENCE column (center)
    int_x = inp_x + col_w + col_gap
    intel_card = add_rect(slide, int_x, Inches(2.0), col_w, Inches(4.5),
                          fill_color=RGBColor(0x0D, 0x2B, 0x1F), border_color=EMERALD)

    add_textbox(slide, int_x, Inches(2.15), col_w, Inches(0.5),
                "🧠", font_size=28, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, int_x, Inches(2.65), col_w, Inches(0.4),
                "INTELLIGENCE ENGINE", font_size=13, font_color=EMERALD_LIGHT,
                bold=True, alignment=PP_ALIGN.CENTER, font_name="Outfit")

    intel_items = [
        "→ FAO-56 Penman-Monteith ET₀",
        "→ Pedotransfer soil hydraulics",
        "→ Root-zone water balance",
        "→ Crop water stress (Kₛ / CWSI)",
        "→ Solar generation modeling",
        "→ Tariff-aware scheduling",
        "→ Traceability chain engine",
    ]
    txBox, tf = add_rich_textbox(slide, int_x + Inches(0.4), Inches(3.15), col_w - Inches(0.8), Inches(3.2))
    for i, item in enumerate(intel_items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(5)
        add_run(p, item, font_size=11, font_color=TEXT_SEC, font_name="Plus Jakarta Sans")

    # OUTPUTS column
    out_x = int_x + col_w + col_gap
    add_textbox(slide, out_x, Inches(1.8), col_w, Inches(0.35),
                "OUTPUTS", font_size=12, font_color=TEXT_TER, bold=True,
                font_name="Outfit")

    outputs = [
        ("🚿", "Irrigation recommendations"),
        ("📊", "Water stress insights (CWSI)"),
        ("⚡", "Solar pumping windows"),
        ("💰", "Financial/energy insights"),
        ("🔗", "Traceable reasoning chain"),
    ]
    for i, (icon, text) in enumerate(outputs):
        y = Inches(2.2) + i * Inches(0.85)
        bc = AMBER_DIM if i == 0 else RGBColor(0x2D, 0x3F, 0x55)
        add_rect(slide, out_x, y, col_w, Inches(0.7), fill_color=BG_CARD, border_color=bc)
        add_textbox(slide, out_x + Inches(0.15), y + Inches(0.12), col_w - Inches(0.3), Inches(0.5),
                    f"{icon}  {text}", font_size=11, font_color=TEXT_SEC, font_name="Plus Jakarta Sans")

    add_footer(slide)
    add_slide_number(slide, 3)


    # ═══════════════════ SLIDE 4: HOW IT THINKS ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "THE PIPELINE")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "From raw observations to an ", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "explainable", font_size=28, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, " irrigation decision.", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")

    # Pipeline image
    img = add_image_safe(slide, "pipeline_architecture.jpg",
                         Inches(0.5), Inches(1.8), width=Inches(12.3))

    # If no image, draw pipeline boxes
    if img is None:
        stages = [
            ("01", "📡", "DATA INGESTION", "Open-Meteo\nISRIC SoilGrids\nSentinel-2\nSHA-256 archived", SKY),
            ("02", "🔬", "AGRI PHYSICS", "FAO-56 Penman-\nMonteith ET₀\nPedotransfer\nθFC, θWP, AWC", EMERALD),
            ("03", "💧", "WATER STRESS", "TAW / RAW\nDepletion Dr\nKₛ coefficient\nCWSI index", EMERALD_LIGHT),
            ("04", "☀️", "ENERGY INTEL", "Solar PV curve\nPump matching\nTariff-aware\n10:30–3:45 PM", AMBER),
            ("05", "✅", "RECOMMENDATION", "Action + Timing\nConfidence\nAssumptions\nLimitations", AMBER_LIGHT),
        ]
        step_w = Inches(2.2)
        arrow_w = Inches(0.3)
        start_x = Inches(0.6)
        for i, (num, icon, title, detail, color) in enumerate(stages):
            x = start_x + i * (step_w + arrow_w)
            card = add_rect(slide, x, Inches(2.0), step_w, Inches(3.8),
                            fill_color=BG_CARD, border_color=color)
            add_textbox(slide, x, Inches(2.15), step_w, Inches(0.3),
                        f"STAGE {num}", font_size=9, font_color=color,
                        alignment=PP_ALIGN.CENTER, font_name="JetBrains Mono")
            add_textbox(slide, x, Inches(2.5), step_w, Inches(0.5),
                        icon, font_size=26, alignment=PP_ALIGN.CENTER)
            add_textbox(slide, x, Inches(3.1), step_w, Inches(0.4),
                        title, font_size=11, font_color=WHITE, bold=True,
                        alignment=PP_ALIGN.CENTER, font_name="Outfit")
            add_textbox(slide, x + Inches(0.15), Inches(3.55), step_w - Inches(0.3), Inches(2.0),
                        detail, font_size=10, font_color=TEXT_TER,
                        alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")
            if i < len(stages) - 1:
                add_textbox(slide, x + step_w, Inches(3.6), arrow_w, Inches(0.4),
                            "→", font_size=18, font_color=EMERALD,
                            alignment=PP_ALIGN.CENTER)

    # Bottom quote
    add_textbox(slide, Inches(1.5), Inches(6.2), Inches(10), Inches(0.4),
                "Every recommendation is traceable: input → calculation → assumption → output.",
                font_size=13, font_color=TEXT_SEC, italic=True,
                alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")

    add_footer(slide)
    add_slide_number(slide, 4)


    # ═══════════════════ SLIDE 5: THE SCIENCE ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "THE SCIENCE")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "Built on ", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "agricultural physics", font_size=28, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, "—not a black box.", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")

    # Left: bullet list
    science_bullets = [
        "FAO-56 Penman-Monteith — deterministic reference evapotranspiration from net radiation, vapor pressure deficit, wind speed, and psychrometric constants.",
        "Saxton-Rawls pedotransfer — converts SoilGrids texture (sand%, clay%) into field capacity (θFC), wilting point (θWP), and available water capacity.",
        "Root-zone water balance — daily depletion mass model tracking Dr against USDA-SCS effective rainfall and adjusted ETc.",
        "Crop water stress — Kₛ coefficient and CWSI derived from depletion vs. readily-available water thresholds.",
        "Effective rainfall — USDA-SCS methodology separating surface runoff from infiltrating precipitation."
    ]
    add_bullet_list(slide, Inches(0.8), Inches(1.8), Inches(6.0), Inches(4.5), science_bullets, font_size=12)

    # Right: equation boxes
    eq_y = Inches(1.8)
    eq_x = Inches(7.2)
    eq_w = Inches(5.5)

    # ET0 equation box
    eq1 = add_rect(slide, eq_x, eq_y, eq_w, Inches(2.0),
                   fill_color=RGBColor(0x0A, 0x10, 0x1E), border_color=EMERALD_DIM)
    add_textbox(slide, eq_x + Inches(0.3), eq_y + Inches(0.15), eq_w - Inches(0.6), Inches(0.3),
                "FAO-56 Penman-Monteith Reference ET₀", font_size=9, font_color=TEXT_TER,
                alignment=PP_ALIGN.CENTER, font_name="JetBrains Mono")
    add_textbox(slide, eq_x + Inches(0.3), eq_y + Inches(0.5), eq_w - Inches(0.6), Inches(0.6),
                "ET₀ = [0.408·Δ·(Rn − G) + γ·(900/(T+273))·u₂·(eₛ − eₐ)]",
                font_size=11, font_color=EMERALD_LIGHT,
                alignment=PP_ALIGN.CENTER, font_name="JetBrains Mono")
    add_textbox(slide, eq_x + Inches(0.3), eq_y + Inches(1.1), eq_w - Inches(0.6), Inches(0.5),
                "/ [Δ + γ·(1 + 0.34·u₂)]",
                font_size=11, font_color=AMBER,
                alignment=PP_ALIGN.CENTER, font_name="JetBrains Mono")

    # Depletion & CWSI equation box
    eq2_y = eq_y + Inches(2.3)
    eq2 = add_rect(slide, eq_x, eq2_y, eq_w, Inches(1.6),
                   fill_color=RGBColor(0x0A, 0x10, 0x1E), border_color=AMBER_DIM)
    add_textbox(slide, eq_x + Inches(0.3), eq2_y + Inches(0.15), eq_w - Inches(0.6), Inches(0.3),
                "Root-Zone Depletion & Stress", font_size=9, font_color=TEXT_TER,
                alignment=PP_ALIGN.CENTER, font_name="JetBrains Mono")
    add_textbox(slide, eq_x + Inches(0.3), eq2_y + Inches(0.5), eq_w - Inches(0.6), Inches(0.4),
                "Dr,i = Dr,i−1 − Peff − I + ETc,adj", font_size=11,
                font_color=EMERALD_LIGHT, alignment=PP_ALIGN.CENTER, font_name="JetBrains Mono")
    add_textbox(slide, eq_x + Inches(0.3), eq2_y + Inches(1.0), eq_w - Inches(0.6), Inches(0.4),
                "CWSI = 1 − Kₛ", font_size=13,
                font_color=AMBER, bold=True, alignment=PP_ALIGN.CENTER, font_name="JetBrains Mono")

    # Traceability note
    add_textbox(slide, eq_x, Inches(6.0), eq_w, Inches(0.5),
                "Scientific calculations remain traceable from\ninput → calculation → assumption → output.",
                font_size=11, font_color=TEXT_SEC, italic=True,
                alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")

    add_footer(slide)
    add_slide_number(slide, 5)


    # ═══════════════════ SLIDE 6: ENERGY INNOVATION ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "THE ENERGY INNOVATION", color=AMBER)

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "Water intelligence meets ", font_size=30, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "solar intelligence", font_size=30, font_color=AMBER, bold=True, font_name="Outfit")
    add_run(p, ".", font_size=30, font_color=WHITE, bold=True, font_name="Outfit")

    # Left: text
    add_textbox(slide, Inches(0.8), Inches(1.8), Inches(5.8), Inches(1.0),
                "Instead of treating irrigation demand and energy cost as separate decisions, "
                "KISAN URJA connects them. The platform models when solar PV generation is "
                "sufficient to power irrigation pumps—synchronizing water delivery with "
                "available renewable energy.",
                font_size=13, font_color=TEXT_SEC, font_name="Plus Jakarta Sans")

    # Decision fusion card
    fusion = add_rect(slide, Inches(0.8), Inches(3.0), Inches(5.5), Inches(3.0),
                      fill_color=RGBColor(0x1A, 0x1A, 0x0A), border_color=AMBER_DIM)
    add_textbox(slide, Inches(1.1), Inches(3.15), Inches(5), Inches(0.3),
                "DECISION FUSION", font_size=11, font_color=AMBER, bold=True, font_name="Outfit")

    fusion_items = [
        "💧  Crop Water Need",
        "  +",
        "🪨  Soil Water Status",
        "  +",
        "☀️  Solar Availability",
        "  +",
        "⚙️  Pump Capacity",
        "  +",
        "💰  Energy Cost / Tariff",
    ]
    txBox, tf = add_rich_textbox(slide, Inches(1.3), Inches(3.5), Inches(4.8), Inches(2.0))
    for i, item in enumerate(fusion_items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(1)
        color = TEXT_TER if item.strip() == "+" else TEXT_SEC
        add_run(p, item, font_size=11, font_color=color, font_name="Plus Jakarta Sans")

    add_textbox(slide, Inches(1.3), Inches(5.5), Inches(4.8), Inches(0.35),
                "→  OPTIMAL IRRIGATION WINDOW", font_size=13, font_color=AMBER,
                bold=True, font_name="Outfit")

    # Right: Solar timeline image
    add_image_safe(slide, "solar_timeline.jpg",
                   Inches(6.8), Inches(1.8), width=Inches(6.0))

    # Disclaimer
    add_textbox(slide, Inches(6.8), Inches(5.7), Inches(6.0), Inches(0.8),
                "Demo profiles model savings scenarios (e.g., ₹94,200/yr for a 5HP solar pump "
                "in Haryana). These represent modeled scenarios, not independently measured field results.",
                font_size=9, font_color=TEXT_TER, italic=True, font_name="Plus Jakarta Sans")

    add_footer(slide)
    add_slide_number(slide, 6)


    # ═══════════════════ SLIDE 7: PRODUCT ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "PRODUCT — LIVE DEMO")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "An intelligence console ", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "built for the farmer", font_size=28, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, ".", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")

    # Left: dashboard mockup
    add_image_safe(slide, "dashboard_mockup.jpg",
                   Inches(0.5), Inches(1.8), width=Inches(7.0))

    # Right: Feature cards
    features = [
        ("👨‍🌾", "Farmer Identity Card", "Full profile with state badge, farm name, solar pump specs, and modeled savings. Multi-state switching.", EMERALD_LIGHT),
        ("💧", "Water Balance & CWSI", "Live root-zone depletion bar, stress index, ET₀, and adjusted crop water demand.", EMERALD_LIGHT),
        ("☀️", "Solar Energy & Pump", "PV generation, pump status (ACTIVE_SOLAR / STANDBY), tariff savings.", AMBER),
        ("⚡", "Live Simulation", "Interactive telemetry: Peak Sun, Rainfall, Heatwave—demonstrating dynamic response.", SKY),
        ("🛰️", "Satellite Scanner", "4-band spectral switcher (NDVI, CWSI, NDRE, True Color) with zone hotspots.", EMERALD_LIGHT),
    ]

    card_x = Inches(7.8)
    card_w = Inches(5.0)
    for i, (icon, title, desc, color) in enumerate(features):
        y = Inches(1.8) + i * Inches(1.0)
        add_rect(slide, card_x, y, card_w, Inches(0.88),
                 fill_color=BG_CARD, border_color=RGBColor(0x2D, 0x3F, 0x55))

        txBox, tf = add_rich_textbox(slide, card_x + Inches(0.15), y + Inches(0.08),
                                     card_w - Inches(0.3), Inches(0.75))
        p = tf.paragraphs[0]
        add_run(p, f"{icon} {title}", font_size=11, font_color=color, bold=True, font_name="Outfit")
        p2 = tf.add_paragraph()
        p2.space_before = Pt(2)
        add_run(p2, desc, font_size=9, font_color=TEXT_SEC, font_name="Plus Jakarta Sans")

    add_textbox(slide, Inches(0.8), Inches(6.85), Inches(5), Inches(0.35),
                "Live: kisanurja.vercel.app", font_size=9, font_color=EMERALD,
                font_name="JetBrains Mono")
    add_textbox(slide, Inches(7), Inches(6.85), Inches(5.5), Inches(0.35),
                "React 19 + Vite 8 · Obsidian Slate Theme · Lenis Smooth Scroll",
                font_size=9, font_color=TEXT_TER, alignment=PP_ALIGN.RIGHT,
                font_name="Plus Jakarta Sans")
    add_slide_number(slide, 7)


    # ═══════════════════ SLIDE 8: ACCESSIBILITY + TRUST ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "ACCESSIBILITY & TRUST")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "Advanced intelligence, ", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "without hiding the reasoning", font_size=28, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, ".", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")

    # Left column: Accessibility
    left_x = Inches(0.8)
    col_w = Inches(5.8)

    add_textbox(slide, left_x, Inches(1.8), col_w, Inches(0.35),
                "🇮🇳  FARMER ACCESSIBILITY", font_size=12, font_color=EMERALD_LIGHT,
                bold=True, font_name="Outfit")

    acc_bullets = [
        "Hindi/Hinglish transliterator — real-time Devanagari conversion with agricultural dictionary (khet → खेत, paani → पानी)",
        "Hindi voice synthesis — Web Speech API reads recommendations aloud",
        "Sunlight high-contrast mode — toggle for outdoor field visibility",
        "Font size scaling — adjustable text size controls (+/−)",
        "Farmer-first interface — clear recommendation cards with urgency badges"
    ]
    add_bullet_list(slide, left_x, Inches(2.2), col_w, Inches(3.5), acc_bullets, font_size=12)

    # Right column: Trust
    right_x = Inches(7.0)

    add_textbox(slide, right_x, Inches(1.8), col_w, Inches(0.35),
                "🔒  TRUST & TRACEABILITY", font_size=12, font_color=AMBER,
                bold=True, font_name="Outfit")

    trust_bullets = [
        "Data provenance — every record tagged EXTERNAL_RETRIEVED, OBSERVED, or CALCULATED",
        "SHA-256 raw archiving — all external API payloads cryptographically hashed before parsing",
        "PostgreSQL Row-Level Security — tenant isolation via JWT claim enforcement",
        "Full traceability chains — each recommendation includes inputs, equations, assumptions, confidence, limitations",
        "Zero fabricated telemetry — no synthetic data presented as observed measurements"
    ]
    add_bullet_list(slide, right_x, Inches(2.2), col_w, Inches(3.5), trust_bullets,
                    font_size=12, bullet_color=AMBER)

    # Bottom quote
    txBox, tf = add_rich_textbox(slide, Inches(1.5), Inches(6.0), Inches(10), Inches(0.6))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, '"Farmers should understand not only ', font_size=14, font_color=TEXT_SEC,
            italic=True, font_name="Plus Jakarta Sans")
    add_run(p, 'what', font_size=14, font_color=WHITE, bold=True, italic=True,
            font_name="Plus Jakarta Sans")
    add_run(p, ' the system recommends, but ', font_size=14, font_color=TEXT_SEC,
            italic=True, font_name="Plus Jakarta Sans")
    add_run(p, 'why', font_size=14, font_color=EMERALD_LIGHT, bold=True, italic=True,
            font_name="Plus Jakarta Sans")
    add_run(p, '."', font_size=14, font_color=TEXT_SEC, italic=True,
            font_name="Plus Jakarta Sans")

    add_footer(slide)
    add_slide_number(slide, 8)


    # ═══════════════════ SLIDE 9: IMPLEMENTATION ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "IMPLEMENTATION")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "From concept to ", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "working full-stack system", font_size=28, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, ".", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")

    # Left: Architecture stack
    add_textbox(slide, Inches(0.8), Inches(1.8), Inches(5), Inches(0.35),
                "SYSTEM ARCHITECTURE", font_size=11, font_color=TEXT_TER,
                bold=True, font_name="Outfit")

    layers = [
        ("🖥️  React 19 + Vite 8 + Lenis", SKY, RGBColor(0x0C, 0x20, 0x30)),
        ("⚡  FastAPI 0.115+ (10 Domain Routers)", EMERALD, RGBColor(0x0D, 0x2B, 0x1F)),
        ("🧠  Agricultural Intelligence Engine", AMBER, RGBColor(0x1A, 0x1A, 0x0A)),
        ("🗄️  PostgreSQL 18 + PostGIS 3.6", TEXT_SEC, BG_CARD),
        ("📡  Open-Meteo · SoilGrids · Sentinel-2", TEXT_TER, BG_CARD),
    ]
    stack_x = Inches(1.2)
    stack_w = Inches(4.8)
    for i, (text, border, bg) in enumerate(layers):
        y = Inches(2.3) + i * Inches(0.82)
        add_rect(slide, stack_x, y, stack_w, Inches(0.55), fill_color=bg, border_color=border)
        add_textbox(slide, stack_x + Inches(0.2), y + Inches(0.08), stack_w - Inches(0.4), Inches(0.4),
                    text, font_size=12, font_color=border, bold=True,
                    alignment=PP_ALIGN.CENTER, font_name="Outfit")
        if i < len(layers) - 1:
            add_textbox(slide, stack_x + Inches(2.1), y + Inches(0.55), Inches(0.6), Inches(0.25),
                        "↕", font_size=12, font_color=TEXT_TER, alignment=PP_ALIGN.CENTER)

    # Right: Stats grid
    add_textbox(slide, Inches(7.0), Inches(1.8), Inches(5.5), Inches(0.35),
                "VERIFIED IMPLEMENTATION", font_size=11, font_color=TEXT_TER,
                bold=True, font_name="Outfit")

    stats = [
        ("77", "Relational Entities", EMERALD_LIGHT),
        ("25", "Verified Migrations", AMBER),
        ("10", "API Domain Routers", EMERALD_LIGHT),
        ("15/15", "Test Suites Passing", AMBER),
    ]
    for i, (val, label, color) in enumerate(stats):
        col = i % 2
        row = i // 2
        x = Inches(7.0) + col * Inches(2.8)
        y = Inches(2.3) + row * Inches(1.5)
        add_rect(slide, x, y, Inches(2.5), Inches(1.2), fill_color=BG_CARD,
                 border_color=RGBColor(0x2D, 0x3F, 0x55))
        add_textbox(slide, x, y + Inches(0.15), Inches(2.5), Inches(0.6),
                    val, font_size=28, font_color=color, bold=True,
                    alignment=PP_ALIGN.CENTER, font_name="Outfit")
        add_textbox(slide, x, y + Inches(0.75), Inches(2.5), Inches(0.3),
                    label, font_size=10, font_color=TEXT_TER,
                    alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")

    # Right: additional bullets
    impl_bullets = [
        "Docker Compose full-stack orchestration",
        "NIST PBKDF2-HMAC-SHA256 auth & JWT",
        "Multi-tenant RLS isolation per farmer",
        "Production deployed on Vercel (frontend)",
        "Pipeline orchestrator with deduplication"
    ]
    add_bullet_list(slide, Inches(7.0), Inches(5.4), Inches(5.5), Inches(1.5),
                    impl_bullets, font_size=11, bold_prefix=False)

    add_textbox(slide, Inches(0.8), Inches(6.85), Inches(6), Inches(0.35),
                "GitHub: nikhillakra2007-tech/yuva-energy", font_size=9,
                font_color=EMERALD, font_name="JetBrains Mono")
    add_textbox(slide, Inches(7), Inches(6.85), Inches(5.5), Inches(0.35),
                "15/15 backend test suites passing — documented in README",
                font_size=9, font_color=TEXT_TER, alignment=PP_ALIGN.RIGHT,
                font_name="Plus Jakarta Sans")
    add_slide_number(slide, 9)


    # ═══════════════════ SLIDE 10: IMPACT + ROADMAP ═══════════════════
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide)

    add_label(slide, Inches(0.8), Inches(0.5), "IMPACT & FUTURE")

    txBox, tf = add_rich_textbox(slide, Inches(0.8), Inches(0.9), Inches(11), Inches(0.7))
    p = tf.paragraphs[0]
    add_run(p, "Toward ", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "autonomous", font_size=28, font_color=EMERALD_LIGHT, bold=True, font_name="Outfit")
    add_run(p, ", ", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "sustainable", font_size=28, font_color=AMBER, bold=True, font_name="Outfit")
    add_run(p, " irrigation.", font_size=28, font_color=WHITE, bold=True, font_name="Outfit")

    # 3 impact cards
    impact_items = [
        ("💧", "Water", "More informed irrigation decisions.\nBetter visibility into root-zone depletion and crop stress.", SKY),
        ("☀️", "Energy", "Greater alignment between irrigation timing and solar availability.\nDesigned to reduce reliance on expensive pumping.", AMBER),
        ("👨‍🌾", "Farmer", "Actionable, explainable recommendations in Hindi/English.\nTraceable decisions that build trust.", EMERALD_LIGHT),
    ]
    for i, (icon, title, desc, color) in enumerate(impact_items):
        x = Inches(0.8) + i * Inches(4.1)
        card = add_rect(slide, x, Inches(1.8), Inches(3.8), Inches(2.0),
                        fill_color=BG_CARD, border_color=color)
        add_textbox(slide, x, Inches(1.9), Inches(3.8), Inches(0.5),
                    icon, font_size=24, alignment=PP_ALIGN.CENTER)
        add_textbox(slide, x, Inches(2.35), Inches(3.8), Inches(0.35),
                    title, font_size=14, font_color=color, bold=True,
                    alignment=PP_ALIGN.CENTER, font_name="Outfit")
        add_textbox(slide, x + Inches(0.3), Inches(2.75), Inches(3.2), Inches(0.9),
                    desc, font_size=10, font_color=TEXT_SEC,
                    alignment=PP_ALIGN.CENTER, font_name="Plus Jakarta Sans")

    # Roadmap: NOW | NEXT
    roadmap_y = Inches(4.1)
    # NOW
    now_card = add_rect(slide, Inches(0.8), roadmap_y, Inches(5.8), Inches(2.0),
                        fill_color=RGBColor(0x0D, 0x2B, 0x1F), border_color=EMERALD)
    add_textbox(slide, Inches(1.1), roadmap_y + Inches(0.12), Inches(5.2), Inches(0.35),
                "✅  NOW — Working Prototype", font_size=13, font_color=EMERALD_LIGHT,
                bold=True, font_name="Outfit")
    now_items = [
        "→ Full-stack platform deployed live",
        "→ FAO-56 physics engine verified",
        "→ 15/15 test suites passing",
        "→ 4-state benchmark profiles",
        "→ Interactive simulation engine"
    ]
    txBox, tf = add_rich_textbox(slide, Inches(1.3), roadmap_y + Inches(0.5), Inches(5.0), Inches(1.4))
    for i, item in enumerate(now_items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(3)
        add_run(p, item, font_size=11, font_color=TEXT_SEC, font_name="Plus Jakarta Sans")

    # NEXT
    next_card = add_rect(slide, Inches(7.0), roadmap_y, Inches(5.8), Inches(2.0),
                         fill_color=RGBColor(0x1A, 0x1A, 0x0A), border_color=AMBER_DIM)
    add_textbox(slide, Inches(7.3), roadmap_y + Inches(0.12), Inches(5.2), Inches(0.35),
                "🔜  NEXT — Validation & Expansion", font_size=13, font_color=AMBER,
                bold=True, font_name="Outfit")
    next_items = [
        "→ Real-world field validation",
        "→ Localized agronomic models",
        "→ Expanded region/crop coverage",
        "→ Stronger live data integration",
        "→ Progressive irrigation autonomy"
    ]
    txBox, tf = add_rich_textbox(slide, Inches(7.5), roadmap_y + Inches(0.5), Inches(5.0), Inches(1.4))
    for i, item in enumerate(next_items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(3)
        add_run(p, item, font_size=11, font_color=TEXT_SEC, font_name="Plus Jakarta Sans")

    # Closing
    add_divider_line(slide, Inches(5.9), Inches(6.3), Inches(1.5))

    txBox, tf = add_rich_textbox(slide, Inches(2), Inches(6.4), Inches(9.3), Inches(0.4))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "KISAN ", font_size=22, font_color=WHITE, bold=True, font_name="Outfit")
    add_run(p, "URJA", font_size=22, font_color=AMBER, bold=True, font_name="Outfit")
    add_run(p, "   ·   ", font_size=14, font_color=TEXT_TER, font_name="Outfit")
    add_run(p, "kisanurja.vercel.app", font_size=12, font_color=EMERALD, font_name="JetBrains Mono")

    add_slide_number(slide, 10)

    return prs


# ──────────────────── MAIN ────────────────────
if __name__ == "__main__":
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                               "KISAN_URJA_Hackathon_2026.pptx")
    prs = build_presentation()
    prs.save(output_path)
    print(f"\n[OK] Presentation saved: {output_path}")
    print(f"     Slides: {len(prs.slides)}")
    print(f"     Format: 13.333 x 7.5 inches (16:9 widescreen)")
